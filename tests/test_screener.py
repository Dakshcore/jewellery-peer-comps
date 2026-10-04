import datetime as dt

import openpyxl

from comps.screener import load_screener, to_crore
from tests.synthetic import TempDirCase


class Loader(TempDirCase):
    def test_reads_annual_series_and_meta(self):
        data = load_screener(self.workbook("ALPHA"))
        self.assertEqual(data.symbol, "ALPHA")
        self.assertEqual(data.name, "Alpha Ltd")
        self.assertEqual(data.annual["sales"], {2023: 1000.0, 2024: 1100.0, 2025: 1210.0, 2026: 1331.0})
        self.assertEqual(data.annual["cash"][2026], 100.0)
        self.assertEqual((data.price, data.market_cap, data.face_value), (100.0, 1000.0, 10.0))

    def test_rows_are_looked_up_inside_their_own_section(self):
        data = load_screener(self.workbook())
        self.assertEqual(len(data.annual["sales"]), 4)       # annual, not the 5 quarters
        self.assertEqual(len(data.quarterly["sales"]), 5)
        self.assertEqual(data.annual["net_profit"][2026], 100.0)
        self.assertEqual(data.quarterly["net_profit"][dt.date(2026, 6, 30)], 28.0)

    def test_labels_ignore_case_spacing_and_colons(self):
        plain = load_screener(self.workbook("A"))
        messy = load_screener(self.workbook("B", style=lambda s: "  " + s.upper().replace(" ", "   ") + " :"))
        self.assertEqual(messy.annual, plain.annual)
        self.assertEqual(messy.quarterly, plain.quarterly)
        self.assertEqual(messy.price, 100.0)

    def test_missing_required_label_is_named_in_the_error(self):
        path = self.workbook(drop={"Profit before tax"})
        with self.assertRaises(ValueError) as caught:
            load_screener(path)
        self.assertIn("Profit before tax", str(caught.exception))
        self.assertIn("PROFIT & LOSS", str(caught.exception))

    def test_missing_section_is_named_in_the_error(self):
        with self.assertRaises(ValueError) as caught:
            load_screener(self.workbook(with_balance_sheet=False))
        self.assertIn("BALANCE SHEET", str(caught.exception))

    def test_missing_data_sheet_and_missing_file(self):
        with self.assertRaises(ValueError) as caught:
            load_screener(self.workbook(sheet_title="Sheet1"))
        self.assertIn("Data Sheet", str(caught.exception))
        with self.assertRaises(FileNotFoundError):
            load_screener(self.dir / "NOPE.xlsx")

    def test_optional_rows_and_quarters_may_be_absent(self):
        data = load_screener(self.workbook(drop={"Inventory"}, quarters=False))
        self.assertNotIn("inventory", data.annual)
        self.assertEqual(data.quarterly, {})

    def test_share_count_is_normalised_to_crore(self):
        self.assertEqual(load_screener(self.workbook("A", shares=1e8)).shares_cr, 10.0)   # absolute number
        self.assertEqual(load_screener(self.workbook("B", shares=10.0)).shares_cr, 10.0)  # already crore
        self.assertAlmostEqual(to_crore(888_000_000), 88.8)

    def test_uncached_shares_formula_reads_as_none_and_shares_come_from_market_cap_over_price(self):
        path = self.workbook("A")                              # default: shares cell is an uncached formula
        sheet = openpyxl.load_workbook(path, data_only=True)["Data Sheet"]
        self.assertEqual(next(r[1] for r in sheet.iter_rows(values_only=True) if r[0] == "Number of shares"), None)
        data = load_screener(path)
        self.assertEqual(data.shares_cr, 10.0)                 # 1000 crore / Rs 100
        self.assertEqual(load_screener(self.workbook("B", mcap=1234.5, price=61.725)).shares_cr, 20.0)

    def test_share_count_fallbacks_when_market_cap_is_missing(self):
        derived = self.workbook("A", shares=None, mcap=None, derived_shares=9.5)
        self.assertEqual(load_screener(derived).shares_cr, 9.5)               # DERIVED Adjusted Equity Shares in Cr
        balance_sheet = self.workbook("B", shares=None, mcap=None, derived_shares=None)
        self.assertEqual(load_screener(balance_sheet).shares_cr, 10.0)        # No. of Equity Shares, absolute
        self.assertIsNone(load_screener(self.workbook("C", shares=None, mcap=None, derived_shares=None,
                                                      balance={"No. of Equity Shares": [None] * 4})).shares_cr)

    def test_cached_shares_value_wins_over_the_derivation(self):
        self.assertEqual(load_screener(self.workbook(shares=2e7, mcap=1000, price=100)).shares_cr, 2.0)

    def test_price_and_derived_rows_do_not_leak_into_the_cash_flow_section(self):
        data = load_screener(self.workbook())
        self.assertEqual((data.price, data.market_cap), (100.0, 1000.0))
        sheet = openpyxl.load_workbook(self.workbook("X"), data_only=True)["Data Sheet"]
        labels = [r[0] for r in sheet.iter_rows(values_only=True)]
        self.assertIn("CASH FLOW:", labels)
        self.assertIn("PRICE:", labels)
        self.assertIn("DERIVED:", labels)

    def test_blank_other_income_cell_is_zero(self):
        data = load_screener(self.workbook())                  # the oldest synthetic quarter has an empty cell
        quarter = data.quarterly["other_income"]
        self.assertEqual(quarter[dt.date(2025, 6, 30)], 0.0)
        self.assertEqual(len(quarter), 5)
        self.assertEqual(quarter[dt.date(2026, 6, 30)], 1.0)
