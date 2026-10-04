import contextlib
import io

import openpyxl

from comps.__main__ import main
from tests.synthetic import TempDirCase

# Same fundamentals everywhere (P/E = price / 10), so the core peers' P/E are 8, 10, 12 and 14.
PRICES = {"TITAN": 100, "KALYANKJIL": 80, "SENCO": 100, "THANGAMAYL": 120, "PNGJL": 140, "TRENT": 1000}


class CommandLine(TempDirCase):
    def setUp(self):
        super().setUp()
        self.screener = self.dir / "screener"
        self.out = self.dir / "out"
        for symbol, price in PRICES.items():
            self.workbook(symbol, folder=self.screener, price=price)

    def run_main(self, *extra):
        argv = ["--data-dir", str(self.screener), "--out-dir", str(self.out), *extra]
        out, err = io.StringIO(), io.StringIO()
        with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
            code = main(argv)
        return code, out.getvalue(), err.getvalue()

    def sheet_rows(self, name):
        workbook = openpyxl.load_workbook(self.out / "comps.xlsx")
        return [list(row) for row in workbook[name].iter_rows(values_only=True)]

    def median_pe(self, header="P/E"):
        rows = self.sheet_rows("Comps table")
        pe_column = rows[0].index(header)
        return next(row[pe_column] for row in rows if row[1] == "Peer median")

    def test_end_to_end_with_screener_prices(self):
        code, out, err = self.run_main()
        self.assertEqual(code, 0)
        self.assertNotIn("skipping", err)
        self.assertIn("TITAN", out)
        workbook = openpyxl.load_workbook(self.out / "comps.xlsx")
        self.assertEqual(workbook.sheetnames, ["Comps table", "Implied values", "Inputs & as-of dates", "Notes"])
        self.assertEqual((self.out / "football_field.png").read_bytes()[:4], b"\x89PNG")
        self.assertAlmostEqual(self.median_pe(), 11.0)        # core peers only; Trent is excluded

    def test_reference_peers_can_be_included_in_the_statistics(self):
        self.run_main("--include-reference")
        self.assertAlmostEqual(self.median_pe(), 12.0)        # P/E 8, 10, 12, 14, 100

    def test_implied_values_sheet_has_my_valuation_and_reference_price(self):
        self.run_main("--valuation", "4000", "--ref-price", "4900")
        rows = self.sheet_rows("Implied values")
        own = next(row for row in rows if row[0] == "My valuation")
        self.assertEqual(own[3], 4000)
        price = next(row for row in rows if row[0] and str(row[0]).startswith("Reference price"))
        self.assertEqual(price[3], 4900)

    def test_reference_price_defaults_to_the_target_price_in_the_data_with_an_as_of_date(self):
        code, out, _ = self.run_main("--as-of", "01-Oct-2026")
        self.assertEqual(code, 0)
        price = next(row for row in self.sheet_rows("Implied values") if row[0] and str(row[0]).startswith("Reference price"))
        self.assertEqual(price[0], "Reference price (01-Oct-2026)")
        self.assertEqual(price[3], 100)                        # TITAN's Screener price, not a hard-coded default
        self.assertIn("Reference price (01-Oct-2026): 100.0", out)
        inputs = {row[0]: row[1] for row in self.sheet_rows("Inputs & as-of dates") if row[0]}
        self.assertEqual(inputs["Prices as of"], "01-Oct-2026")

    def test_reference_price_override_on_the_command_line(self):
        self.run_main("--ref-price", "4900", "--ref-date", "02-Oct-2026")
        price = next(row for row in self.sheet_rows("Implied values") if row[0] and str(row[0]).startswith("Reference price"))
        self.assertEqual((price[0], price[3]), ("Reference price (02-Oct-2026)", 4900))

    def test_bhavcopy_prices_replace_screener_prices(self):
        bhavcopy = self.dir / "bhav.csv"
        lines = ["TradDt,TckrSymb,SctySrs,ClsPric"] + [f"2026-09-25,{s},EQ,{p * 2}" for s, p in PRICES.items()]
        bhavcopy.write_text("\n".join(lines) + "\n")
        code, _, err = self.run_main("--price-source", "bhavcopy", "--bhavcopy", str(bhavcopy))
        self.assertEqual(code, 0)
        self.assertNotIn("skipping", err)
        rows = self.sheet_rows("Comps table")
        self.assertEqual(rows[1][rows[0].index("Price (₹)")], 200)   # TITAN: 100 in Screener, 200 in bhavcopy
        self.assertAlmostEqual(self.median_pe(), 22.0)

    def test_ltm_basis_runs(self):
        code, _, _ = self.run_main("--basis", "ltm")
        self.assertEqual(code, 0)
        self.assertAlmostEqual(self.median_pe("LTM P/E"), 1100 / 103)     # peer LTM P/E = price x 10 / 103

    def test_missing_target_workbook_is_an_error(self):
        (self.screener / "TITAN.xlsx").unlink()
        code, _, err = self.run_main()
        self.assertEqual(code, 2)
        self.assertIn("TITAN", err)

    def test_missing_peer_workbook_is_skipped_with_a_warning(self):
        (self.screener / "SENCO.xlsx").unlink()
        code, _, err = self.run_main()
        self.assertEqual(code, 0)
        self.assertIn("SENCO", err)
        skipped = next(row for row in self.sheet_rows("Inputs & as-of dates") if row[0] == "Skipped (could not load)")
        self.assertEqual(skipped[1], "SENCO")
