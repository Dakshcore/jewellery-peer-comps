"""Write outputs/comps.xlsx: Comps table, Implied values, Inputs & as-of dates, Notes. Values, not formulas."""
from __future__ import annotations

from collections.abc import Sequence
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .metrics import TABLE_COLUMNS, CompanyMetrics
from .valuation import MULTIPLES, ImpliedRange, PeerStats

NUMBER_FORMATS = {"text": "General", "price": "#,##0.00", "crore": "#,##0", "x": '0.0"x"', "pct": "0.0%",
                  "days": "0"}
STAT_ROWS = (("Peer minimum", "minimum"), ("Peer 25th percentile", "p25"), ("Peer median", "median"),
             ("Peer 75th percentile", "p75"), ("Peer maximum", "maximum"))
HEADER_FILL = PatternFill("solid", fgColor="1F3864")


def _header(ws: Worksheet, headers: Sequence[str], row: int = 1) -> None:
    for column, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=column, value=text)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _widths(ws: Worksheet, minimum: int = 10, maximum: int = 60) -> None:
    for column in ws.columns:
        longest = max((len(str(cell.value)) for cell in column if cell.value is not None), default=0)
        ws.column_dimensions[get_column_letter(column[0].column)].width = min(max(longest + 2, minimum), maximum)


def _comps_sheet(ws: Worksheet, target: CompanyMetrics, peers: Sequence[CompanyMetrics],
                 reference: Sequence[CompanyMetrics], stats: dict[str, PeerStats], basis: str) -> None:
    ws.title = "Comps table"
    _header(ws, ["Group"] + [header for header, _, _ in TABLE_COLUMNS])
    column_of = {attr: index for index, (_, attr, _) in enumerate(TABLE_COLUMNS, start=2)}

    def write_company(row: int, company: CompanyMetrics, group: str) -> None:
        ws.cell(row=row, column=1, value=group)
        for column, (_, attr, kind) in enumerate(TABLE_COLUMNS, start=2):
            value = getattr(company, attr)
            cell = ws.cell(row=row, column=column, value="n/a" if value is None else value)
            cell.number_format = NUMBER_FORMATS[kind]
            if value is None:
                cell.alignment = Alignment(horizontal="right")

    row = 2
    for company, group in [(target, "Target")] + [(p, "Core peer") for p in peers]:
        write_company(row, company, group)
        row += 1
    row += 1
    for label, field in STAT_ROWS:
        ws.cell(row=row, column=1, value="Peer statistic").font = Font(bold=True)
        ws.cell(row=row, column=2, value=label).font = Font(bold=True)
        for key in MULTIPLES:
            attr = f"ltm_{key}" if basis == "ltm" else key
            value = getattr(stats[key], field)
            cell = ws.cell(row=row, column=column_of[attr], value="n/a" if value is None else value)
            cell.number_format = NUMBER_FORMATS["x"]
            cell.font = Font(bold=True)
        row += 1
    if reference:
        row += 1
        for company in reference:
            write_company(row, company, "Reference (lifestyle)")
            row += 1
    ws.freeze_panes = "C2"
    _widths(ws)


def _implied_sheet(ws: Worksheet, ranges: Sequence[ImpliedRange], reference_price: float, reference_date: str) -> None:
    ws.title = "Implied values"
    _header(ws, ["Method", "Minimum (₹/share)", "25th percentile (₹/share)", "Median (₹/share)",
                 "75th percentile (₹/share)", "Maximum (₹/share)", "Median vs reference price"])
    for row, r in enumerate(ranges, start=2):
        ws.cell(row=row, column=1, value=r.label)
        for column, value in enumerate((r.minimum, r.low, r.mid, r.high, r.maximum), start=2):
            ws.cell(row=row, column=column, value=value).number_format = "#,##0"
        ws.cell(row=row, column=7, value=r.mid / reference_price - 1).number_format = "0.0%"
    row = len(ranges) + 3
    ws.cell(row=row, column=1, value=f"Reference price ({reference_date})").font = Font(bold=True)
    ws.cell(row=row, column=4, value=reference_price).number_format = "#,##0"
    _widths(ws)


def _inputs_sheet(ws: Worksheet, inputs: Sequence[tuple[str, str]], companies: Sequence[CompanyMetrics]) -> None:
    ws.title = "Inputs & as-of dates"
    _header(ws, ["Input", "Value"])
    row = 2
    for label, value in inputs:
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=value)
        row += 1
    row += 1
    _header(ws, ["Symbol", "Price used (₹)", "Fiscal year", "LTM quarter end"], row=row)
    for company in companies:
        row += 1
        ws.cell(row=row, column=1, value=company.symbol)
        ws.cell(row=row, column=2, value=company.price).number_format = "#,##0.00"
        ws.cell(row=row, column=3, value=f"FY{company.fy}")
        ws.cell(row=row, column=4, value=company.ltm_end.isoformat() if company.ltm_end else "n/a")
    _widths(ws)


def _notes_sheet(ws: Worksheet, notes: Sequence[str]) -> None:
    ws.title = "Notes"
    _header(ws, ["Notes"])
    for row, note in enumerate(notes, start=2):
        cell = ws.cell(row=row, column=1, value=note)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 120


def write_workbook(path: str | Path, target: CompanyMetrics, peers: Sequence[CompanyMetrics],
                   reference: Sequence[CompanyMetrics], stats: dict[str, PeerStats],
                   ranges: Sequence[ImpliedRange], reference_price: float, reference_date: str,
                   inputs: Sequence[tuple[str, str]], notes: Sequence[str], basis: str = "fy") -> Path:
    """Write the four-sheet workbook and return its path."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    _comps_sheet(workbook.active, target, peers, reference, stats, basis)
    _implied_sheet(workbook.create_sheet(), ranges, reference_price, reference_date)
    _inputs_sheet(workbook.create_sheet(), inputs, [target, *peers, *reference])
    _notes_sheet(workbook.create_sheet(), notes)
    workbook.save(path)
    return path
